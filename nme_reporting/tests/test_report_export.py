import io
from datetime import timedelta

from openpyxl import load_workbook

from odoo.exceptions import UserError
from odoo.tests import tagged

from odoo.addons.nme_invoicing.tests.common import NmeInvoicingCommon


@tagged("post_install", "-at_install")
class TestNmeReportExport(NmeInvoicingCommon):
    def _export(self, report_type, **values):
        wizard = self.env["nme.report.export"].create(
            {"report_type": report_type, **values}
        )
        action = wizard.action_export_xlsx()
        self.assertIn(f"/web/content/nme.report.export/{wizard.id}/", action["url"])
        workbook = load_workbook(io.BytesIO(bytes(wizard.file_data)))
        sheet = workbook.active
        return wizard, [
            [cell.value for cell in row] for row in sheet.iter_rows(min_row=4)
        ]

    def _overdue_invoice(self):
        invoice = self._create_invoice(
            invoice_date=self._today() - timedelta(days=50),
            invoice_date_due=self._today() - timedelta(days=35),
        )
        invoice.action_post()
        return invoice

    def test_stock_export(self):
        wizard, rows = self._export("stock", warehouse_id=self.warehouse.id)
        self.assertEqual(wizard.file_name, "Stock by Warehouse.xlsx")
        self.assertEqual(rows[0][0], "Warehouse")
        product_row = next(row for row in rows[1:] if row[2] == "[TST-1] Test Rice")
        self.assertEqual(product_row[3], 100)
        self.assertEqual(product_row[6], 100 * 80000)

    def test_aging_export(self):
        invoice = self._overdue_invoice()
        _wizard, rows = self._export("receivable")
        invoice_row = next(row for row in rows[1:] if row[1] == invoice.name)
        self.assertEqual(invoice_row[3], 35)
        self.assertEqual(invoice_row[4], "31-60 Days")
        self.assertEqual(invoice_row[5], invoice.amount_residual)

    def test_sales_export(self):
        order = self._confirmed_delivered_order(qty=2)
        _wizard, rows = self._export("sales")
        product_row = next(row for row in rows[1:] if row[0] == "[TST-1] Test Rice")
        self.assertGreaterEqual(product_row[1], 2)
        self.assertGreaterEqual(product_row[3], order.amount_untaxed)

    def test_aging_pdf(self):
        invoice = self._overdue_invoice()
        wizard = self.env["nme.report.export"].create({"report_type": "receivable"})
        action = wizard.action_print_pdf()
        self.assertEqual(action["report_name"], "nme_reporting.report_nme_aging")
        html, _format = self.env["ir.actions.report"]._render_qweb_html(
            "nme_reporting.report_nme_aging", wizard.ids
        )
        content = html.decode()
        self.assertIn("Receivable Aging", content)
        self.assertIn(invoice.partner_id.name, content)

        stock_wizard = self.env["nme.report.export"].create({"report_type": "stock"})
        with self.assertRaises(UserError):
            stock_wizard.action_print_pdf()
